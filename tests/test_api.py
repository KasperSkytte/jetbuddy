import io

import openpyxl

from app import db
from tests.conftest import PASSWORD, new_budget


def member(state, name):
    return next(m for m in state["members"] if m["name"] == name)


def account(state, kind):
    return next(a for a in state["accounts"] if a["kind"] == kind)


def add(client, state, who, kind, name, amounts, year=2027, **extra):
    r = client.post(
        "/api/items",
        json={"member_id": member(state, who)["id"], "kind": kind, "name": name, "year": year, "amounts": amounts, **extra},
    )
    assert r.status_code == 200, r.text
    return r.json()


def test_setup_creates_members_shared_categories_accounts(client, household):
    names = [m["name"] for m in household["members"]]
    assert names == ["Alex", "Sam", "Kim", "Robin", "Shared"]
    assert member(household, "Shared")["is_shared"]
    assert {c["kind"] for c in household["categories"]} == {"income", "expense"}
    assert sorted(a["kind"] for a in household["accounts"]) == ["main", "savings"]


def test_amounts_are_exact_and_summed(client, household):
    add(client, household, "Alex", "income", "Salary", [31250] * 12)
    add(client, household, "Alex", "expense", "Phone", [149.95] * 12)
    add(client, household, "Shared", "expense", "Rent", [12000] * 12)
    s = client.get("/api/summary/2027").json()
    assert s["income"] == [31250] * 12
    assert s["expense"][0] == 12149.95
    assert round(sum(s["expense"]), 2) == 145799.4  # 149.95 * 12 + 12000 * 12
    assert s["net"][0] == 31250 - 12149.95
    assert s["cumulative"][-1] == round((31250 - 12149.95) * 12, 2)


def test_single_cell_edit_and_clear(client, household):
    item = add(client, household, "Sam", "expense", "Dentist", [None] * 12)
    assert client.put(f"/api/items/{item['id']}/amount", json={"year": 2027, "month": 4, "amount": 650}).status_code == 200
    items = client.get("/api/years/2027/items").json()
    assert items[0]["amounts"][3] == 650
    client.put(f"/api/items/{item['id']}/amount", json={"year": 2027, "month": 4, "amount": None})
    items = client.get("/api/years/2027/items").json()
    assert items[0]["amounts"] == [None] * 12  # still in the year, just empty


def test_accounts_follow_budget_and_typed_in_balances(client, household):
    savings = account(household, "savings")
    main = account(household, "main")
    add(client, household, "Alex", "income", "Salary", [1000] * 12)
    add(client, household, "Shared", "expense", "Savings", [300] * 12, account_id=savings["id"])
    client.put(f"/api/accounts/{savings['id']}/balance", json={"year": 2027, "month": 1, "amount": 87000})
    client.put(f"/api/accounts/{main['id']}/balance", json={"year": 2027, "month": 1, "amount": 7127})
    # A real balance later in the year overrides the computed one from then on.
    client.put(f"/api/accounts/{main['id']}/balance", json={"year": 2027, "month": 6, "amount": 5000})
    s = client.get("/api/summary/2027").json()
    acc = {a["id"]: a for a in s["accounts"]}
    assert acc[savings["id"]]["balances"][1] == 87300
    assert acc[savings["id"]]["balances"][11] == 87000 + 300 * 11
    assert acc[main["id"]]["balances"][1] == 7127 + 700
    assert acc[main["id"]]["balances"][5] == 5000
    assert acc[main["id"]]["manual"][5] is True
    assert acc[main["id"]]["balances"][6] == 5700
    assert s["savings"] == [300] * 12
    # The balance carries over into next year.
    s2 = client.get("/api/summary/2028").json()
    acc2 = {a["id"]: a for a in s2["accounts"]}
    assert acc2[savings["id"]]["balances"][0] == 87000 + 300 * 11


def test_copy_year_and_remove_from_year(client, household):
    rent = add(client, household, "Shared", "expense", "Rent", [10000] * 12)
    add(client, household, "Shared", "expense", "Holiday", [None] * 6 + [2500] + [None] * 5)
    assert client.get("/api/years/2028/items").json() == []
    r = client.post("/api/years/2028/copy", json={"from_year": 2027})
    assert r.json() == {"copied": 2}
    items = client.get("/api/years/2028/items").json()
    assert len(items) == 2
    assert items[0]["prev"] == [10000] * 12
    # Copying again adds nothing.
    assert client.post("/api/years/2028/copy", json={"from_year": 2027}).json() == {"copied": 0}
    # Taking it out of 2028 leaves 2027 alone.
    client.delete(f"/api/items/{rent['id']}", params={"year": 2028})
    assert len(client.get("/api/years/2028/items").json()) == 1
    assert len(client.get("/api/years/2027/items").json()) == 2
    # Taking it out of its last year deletes it.
    client.delete(f"/api/items/{rent['id']}", params={"year": 2027})
    assert len(client.get("/api/years/2027/items").json()) == 1


def test_trends(client, household):
    add(client, household, "Alex", "income", "Salary", [1000] * 12, year=2026)
    add(client, household, "Alex", "income", "Salary 2", [1100] * 12, year=2027)
    t = client.get("/api/trends").json()
    assert [y["year"] for y in t] == [2026, 2027]
    assert t[1]["income"] == 13200


def test_backup_roundtrip(client, household):
    add(client, household, "Alex", "income", "Salary", [1000.5] * 12)
    backup = client.get("/api/export").json()
    client.delete(f"/api/items/{backup['items'][0]['id']}")
    assert client.get("/api/years/2027/items").json() == []
    import json

    assert client.post("/api/restore", content=json.dumps(backup)).status_code == 200
    assert client.get("/api/years/2027/items").json()[0]["amounts"] == [1000.5] * 12


def test_import_spreadsheet_merges_items_across_years(client, household):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alex ud"
    ws.append(["", "Jan", "Feb", "Mar", "Apr", "Maj", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dec"])
    ws.append(["Phone", *["-149,95"] * 12])
    ws.append(["Dentist", None, None, None, -650] + [None] * 8)
    summary = wb.create_sheet("Summary")
    summary.append(["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
    summary.append(["Total", *[f"=SUM('Alex ud'!{c}:{c})" for c in "BCDEFGHIJKLM"]])
    buf = io.BytesIO()
    wb.save(buf)

    r = client.post("/api/import/preview", files={"file": ("budget.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    sheets = {s["name"]: s for s in r.json()["sheets"]}
    alex = sheets["Alex ud"]
    assert alex["kind"] == "expense"
    assert alex["rows"][0]["amounts"] == [149.95] * 12
    assert alex["rows"][1]["amounts"][3] == 650
    assert sheets["Summary"]["looks_like_summary"]

    alex_id = member(household, "Alex")["id"]
    for year in (2026, 2027):
        r = client.post(
            "/api/import/commit",
            json={"year": year, "sheets": [{"member_id": alex_id, "kind": "expense", "rows": alex["rows"]}]},
        )
        assert r.status_code == 200, r.text
    assert r.json() == {"created": 0, "updated": 2, "suspicious": 0}
    items = client.get("/api/years/2027/items").json()
    assert len(items) == 2
    assert items[0]["prev"] == [149.95] * 12


def test_login_required(client):
    assert client.get("/api/session").json() == {"logged_in": False, "budget": None, "has_budgets": False, "can_create": True}
    assert client.get("/api/state").status_code == 401


def test_login_and_logout(client, household):
    assert client.get("/api/session").json()["logged_in"]  # creating a budget logs in
    client.post("/api/logout")
    assert client.get("/api/state").status_code == 401
    assert client.post("/api/login", json={"login": "test", "password": "nope"}).status_code == 401
    assert client.post("/api/login", json={"login": "nobody", "password": PASSWORD}).status_code == 401
    assert client.post("/api/login", json={"login": " TEST ", "password": PASSWORD}).status_code == 200
    assert client.get("/api/state").json()["settings"]["household"] == "Test"


def test_too_many_wrong_passwords(client, household, monkeypatch):
    monkeypatch.setattr("app.main.time.sleep", lambda s: None)
    for _ in range(10):
        client.post("/api/login", json={"login": "test", "password": "nope"})
    assert client.post("/api/login", json={"login": "test", "password": PASSWORD}).status_code == 429


def test_budgets_are_separate(client, household):
    add(client, household, "Alex", "income", "Salary", [1000] * 12)
    client.post("/api/logout")
    other = new_budget(client, login="other", members=["Pat"])
    assert other.status_code == 200
    assert client.get("/api/years/2027/items").json() == []
    assert new_budget(client, login="OTHER").status_code == 409  # username taken, case-insensitively
    client.post("/api/login", json={"login": "test", "password": PASSWORD})
    assert len(client.get("/api/years/2027/items").json()) == 1


def test_new_budgets_can_be_turned_off(client, household, monkeypatch):
    monkeypatch.setenv("SPAREGRISEN_ALLOW_NEW_BUDGETS", "false")
    assert client.get("/api/session").json()["can_create"] is False
    assert new_budget(client, login="other").status_code == 403


def test_change_password_logs_out_other_devices(client, household):
    other_device = client.cookies.get("sparegrisen_session")
    r = client.put("/api/login-details", json={"current_password": PASSWORD, "password": "new password!", "login": "hansen"})
    assert r.status_code == 200
    assert client.get("/api/state").status_code == 200  # this device stays in
    client.cookies.set("sparegrisen_session", other_device)
    assert client.get("/api/state").status_code == 401
    assert client.post("/api/login", json={"login": "hansen", "password": "new password!"}).status_code == 200
    bad = client.put("/api/login-details", json={"current_password": "wrong", "password": "whatever123"})
    assert bad.status_code == 401


def test_delete_budget(client, household, tmp_path):
    budget_id = client.get("/api/session").json()["budget"]["id"]
    assert (tmp_path / "budgets" / str(budget_id)).exists()
    assert client.post("/api/budget/delete", json={"password": "wrong"}).status_code == 401
    assert client.post("/api/budget/delete", json={"password": PASSWORD}).status_code == 200
    assert not (tmp_path / "budgets" / str(budget_id)).exists()
    assert client.get("/api/session").json()["has_budgets"] is False


def test_short_password_rejected(client):
    r = client.post("/api/budgets", json={"household": "x", "members": ["a"], "login": "x1", "password": "short"})
    assert r.status_code == 422


def test_excel_export(client, household):
    savings = account(household, "savings")
    cats = {c["name"]: c["id"] for c in household["categories"]}
    add(client, household, "Alex", "income", "Salary", [31250] * 12)
    add(client, household, "Sam", "income", "Salary", [22000] * 12)
    add(client, household, "Shared", "expense", "Rent", [12000] * 12, category_id=cats["Housing"], fixed=True)
    add(client, household, "Shared", "expense", "Groceries", [8500] * 12, category_id=cats["Groceries"])
    add(client, household, "Alex", "expense", "Phone", [149.95] * 12, category_id=cats["Subscriptions"], fixed=True)
    add(client, household, "Shared", "expense", "Savings", [9000] * 12, category_id=cats["Savings"], account_id=savings["id"])
    r = client.get("/api/export/2027.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "budget-2027.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:2] == ["Summary", "Budget"]
    summary = {row[0]: row[1:] for row in wb["Summary"].iter_rows(values_only=True) if row[0]}
    assert summary["Total income"][:2] == (53250, 639000)
    # Expenses are negative, as in the app.
    assert summary["Total expenses"][1] == -round(12000 * 12 + 8500 * 12 + 149.95 * 12 + 9000 * 12, 2)
    assert summary["   of which fixed expenses"][1] == -round((12000 + 149.95) * 12, 2)
    budget = [row for row in wb["Budget"].iter_rows(values_only=True)]
    labels = [row[0] for row in budget]
    # Expenses are categorised, so grouped by category in the categories' order.
    assert labels.index("Housing") < labels.index("Groceries") < labels.index("Subscriptions")
    surplus = next(row for row in budget if row[0] == "Surplus (income − expenses)")
    assert surplus[3] == round(53250 - (12000 + 8500 + 149.95 + 9000), 2)
    rent = next(row for row in budget if row[0] == "Rent")
    assert rent[1] == "Shared" and rent[2] == "✓" and rent[3] == -12000
    salary = next(row for row in budget if row[0] == "Salary")
    assert salary[3] > 0
    total_expenses = next(row for row in budget if row[0] == "Total expenses")
    assert total_expenses[3] == -round(12000 + 8500 + 149.95 + 9000, 2)
    # Grouped by person when asked.
    wb2 = openpyxl.load_workbook(io.BytesIO(client.get("/api/export/2027.xlsx?group=member").content))
    labels2 = [row[0] for row in wb2["Budget"].iter_rows(values_only=True)]
    assert "Housing" not in labels2 and labels2.count("Shared") >= 1


def test_import_in_thousands(client, household):
    alex = member(household, "Alex")["id"]
    r = client.post("/api/import/commit", json={
        "year": 2027, "scale": 1000,
        "sheets": [{"member_id": alex, "kind": "income", "rows": [{"name": "Salary", "amounts": [27.8] + [None] * 11}]}],
    })
    assert r.status_code == 200, r.text
    assert client.get("/api/years/2027/items").json()[0]["amounts"][0] == 27800


def test_checks_find_and_fix_amounts_off_by_powers_of_ten(client, household):
    cats = {c["name"]: c["id"] for c in household["categories"]}
    # Salary with one month imported 100 times too big.
    salary = add(client, household, "Alex", "income", "Salary", [31250] * 5 + [3125000] + [31250] * 6)
    # Phone: 2026 right, 2027 1000 times too small (149,95 read as 0,14995). Same
    # number of months each way, so either year could be the wrong one.
    phone = add(client, household, "Alex", "expense", "Phone", [149.95] * 12, year=2026)
    client.put(f"/api/items/{phone['id']}/amounts", json={"year": 2027, "amounts": [0.18] * 12})
    # Rent wrong in every month: only the comparison with income can see it.
    add(client, household, "Shared", "expense", "Rent", [1200000] * 12, category_id=cats["Housing"])
    found = client.get("/api/checks").json()["findings"]
    salary_f = next(f for f in found if f["name"] == "Salary")
    assert salary_f["type"] == "item" and salary_f["where"] == [{"year": 2027, "months": [6]}] and salary_f["factor"] == 100
    assert salary_f["options"] == [{"action": "scale", "multiply": 0.01, "where": [{"year": 2027, "months": [6]}]}]
    phone_f = next(f for f in found if f["name"] == "Phone")
    assert phone_f["type"] == "split" and phone_f["factor"] == 1000
    assert [(o["multiply"], o["where"][0]["year"]) for o in phone_f["options"]] == [(0.001, 2026), (1000, 2027)]
    assert sum(1 for f in found if f["name"] == "Rent" and f["type"] == "large") == 12

    # Fix the salary month; fix the phone's 2027; say the rent in January is right.
    opt = salary_f["options"][0]
    r = client.post("/api/checks/fix", json={"item_id": salary["id"], "where": opt["where"], "multiply": opt["multiply"]})
    assert r.status_code == 200, r.text
    opt = phone_f["options"][1]
    client.post("/api/checks/fix", json={"item_id": phone["id"], "where": opt["where"], "multiply": opt["multiply"]})
    items = {i["name"]: i for i in client.get("/api/years/2027/items").json()}
    assert items["Salary"]["amounts"][5] == 31250 and items["Phone"]["amounts"][0] == 180
    rent_jan = next(f for f in client.get("/api/checks").json()["findings"] if f["name"] == "Rent")
    client.post("/api/checks/dismiss", json={"id": rent_jan["id"]})
    left = client.get("/api/checks").json()["findings"]
    assert rent_jan["id"] not in {f["id"] for f in left}
    assert not any(f["name"] in ("Salary", "Phone") for f in left)
    client.post("/api/checks/reset")
    assert any(f["id"] == rent_jan["id"] for f in client.get("/api/checks").json()["findings"])


def test_checks_typed_in_balance(client, household):
    main = account(household, "main")
    add(client, household, "Alex", "income", "Salary", [30000] * 12)
    add(client, household, "Shared", "expense", "Rent", [20000] * 12)
    client.put(f"/api/accounts/{main['id']}/balance", json={"year": 2027, "month": 1, "amount": 50000})
    client.put(f"/api/accounts/{main['id']}/balance", json={"year": 2027, "month": 3, "amount": 7000000})  # meant 70.000
    f = next(f for f in client.get("/api/checks").json()["findings"] if f["type"] == "balance")
    assert f["where"] == [{"year": 2027, "months": [3]}] and f["factor"] == 100
    opt = f["options"][0]
    client.post("/api/checks/fix", json={"account_id": main["id"], "where": opt["where"], "multiply": opt["multiply"]})
    assert client.get(f"/api/accounts/{main['id']}/balances").json()[1]["amount"] == 70000


def test_import_reports_suspicious_amounts(client, household):
    alex = member(household, "Alex")["id"]
    r = client.post("/api/import/commit", json={"year": 2027, "sheets": [{"member_id": alex, "kind": "income", "rows": [
        {"name": "Salary", "amounts": [31250] * 11 + [3125000]}]}]})
    assert r.json()["suspicious"] == 1


def test_unreadable_import_cells_stay_visible_until_fixed(client, household):
    alex = member(household, "Alex")["id"]
    r = client.post("/api/import/commit", json={"year": 2027, "sheets": [{"member_id": alex, "kind": "expense", "rows": [
        {"name": "Phone", "amounts": [149.95] * 10 + [None, None], "invalid": {"11": "1.2,3.4", "12": "14,9.95"}}]}]})
    assert r.json()["suspicious"] == 1
    item = client.get("/api/years/2027/items").json()[0]
    assert item["unreadable"][10:] == ["1.2,3.4", "14,9.95"]
    f = next(f for f in client.get("/api/checks").json()["findings"] if f["type"] == "unreadable")
    assert f["where"] == [{"year": 2027, "months": [11, 12]}] and f["raw"] == ["1.2,3.4", "14,9.95"]
    # Typing the amount for November settles that cell.
    client.put(f"/api/items/{item['id']}/amount", json={"year": 2027, "month": 11, "amount": 149.95})
    f = next(f for f in client.get("/api/checks").json()["findings"] if f["type"] == "unreadable")
    assert f["where"] == [{"year": 2027, "months": [12]}]
    # Ignoring December leaves it empty and stops reporting it.
    client.post("/api/checks/dismiss", json={"id": f["id"], "item_id": item["id"], "where": f["where"]})
    assert not any(f["type"] == "unreadable" for f in client.get("/api/checks").json()["findings"])
    assert client.get("/api/years/2027/items").json()[0]["unreadable"] == [None] * 12


def test_budget_name_is_the_login(client):
    r = client.post("/api/budgets", json={"household": "Familien Hansen", "members": ["A"], "password": PASSWORD})
    assert r.status_code == 200, r.text
    assert r.json()["accounts"] == []  # accounts come from the bank or are added by hand
    client.post("/api/logout")
    # Any case and spacing.
    assert client.post("/api/login", json={"login": "  familien   HANSEN ", "password": PASSWORD}).status_code == 200
    # Renaming the budget renames the login.
    client.put("/api/settings", json={"household": "Hansen-Berg"})
    client.post("/api/logout")
    assert client.post("/api/login", json={"login": "Familien Hansen", "password": PASSWORD}).status_code == 401
    assert client.post("/api/login", json={"login": "hansen-berg", "password": PASSWORD}).status_code == 200
    # Another budget can't take the name.
    client.post("/api/logout")
    new_budget(client, login=None, household="Other")
    assert client.put("/api/settings", json={"household": "HANSEN-BERG"}).status_code == 409


def test_decimals_rounded_as_set(client, household):
    phone = add(client, household, "Alex", "expense", "Phone", [149.95] * 12)
    refund = add(client, household, "Alex", "expense", "Refund", [-12.3] + [None] * 11)
    for mode, typed, stored in [("nearest", 178.5, 179), ("nearest", 178.49, 178), ("up", 178.01, 179), ("down", 178.99, 178)]:
        assert client.put("/api/settings", json={"decimals": mode}).status_code == 200
        client.put(f"/api/items/{phone['id']}/amount", json={"year": 2027, "month": 1, "amount": typed})
        assert client.get("/api/years/2027/items").json()[0]["amounts"][0] == stored, mode
    # Up and down are away from and towards zero.
    client.put("/api/settings", json={"decimals": "up"})
    client.put(f"/api/items/{refund['id']}/amounts", json={"year": 2027, "amounts": [-12.3] + [None] * 11})
    assert client.get("/api/years/2027/items").json()[1]["amounts"][0] == -13
    # Amounts from before still have decimals: the check finds them and rounds them.
    f = next(f for f in client.get("/api/checks").json()["findings"] if f["type"] == "decimals")
    assert f["item_id"] == phone["id"] and f["where"] == [{"year": 2027, "months": list(range(2, 13))}]
    assert [o["mode"] for o in f["options"]] == ["up"]
    o = f["options"][0]
    client.post("/api/checks/fix", json={"item_id": phone["id"], "where": o["where"], "round": o["mode"]})
    assert client.get("/api/years/2027/items").json()[0]["amounts"][1:] == [150] * 11
    assert not any(f["type"] == "decimals" for f in client.get("/api/checks").json()["findings"])


def test_whole_numbers_only(client, household):
    phone = add(client, household, "Alex", "expense", "Phone", [149.95] + [None] * 11)
    client.put("/api/settings", json={"decimals": "only"})
    r = client.put(f"/api/items/{phone['id']}/amount", json={"year": 2027, "month": 2, "amount": 149.95})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "WHOLE_NUMBERS"
    assert client.put(f"/api/items/{phone['id']}/amount", json={"year": 2027, "month": 2, "amount": 179}).status_code == 200
    # The old amount with decimals is reported, with a choice of how to round it.
    f = next(f for f in client.get("/api/checks").json()["findings"] if f["type"] == "decimals")
    assert f["values"] == [149.95] and [o["mode"] for o in f["options"]] == ["nearest", "up", "down"]
    # Importing: anything but the digits 0-9 can't be read.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alex ud"
    ws.append(["", "Jan", "Feb", "Mar", "Apr", "Maj", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dec"])
    ws.append(["Phone", "179", "149,95", "1.179", 179, 149.95, "179 kr"] + [None] * 6)
    buf = io.BytesIO()
    wb.save(buf)
    row = client.post("/api/import/preview", files={"file": ("b.xlsx", buf.getvalue())}).json()["sheets"][0]["rows"][0]
    assert row["amounts"][:6] == [179, None, None, 179, None, None]
    assert row["invalid"] == {"2": "149,95", "3": "1.179", "5": "149.95", "6": "179 kr"}


def test_import_rounds_decimals(client, household):
    client.put("/api/settings", json={"decimals": "down"})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alex ud"
    ws.append(["", "Jan", "Feb", "Mar", "Apr", "Maj", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dec"])
    ws.append(["Phone", "-149,95", -99.9] + [None] * 10)
    buf = io.BytesIO()
    wb.save(buf)
    row = client.post("/api/import/preview", files={"file": ("b.xlsx", buf.getvalue())}).json()["sheets"][0]["rows"][0]
    assert row["amounts"][:2] == [149, 99]  # flipped to positive, then the decimals dropped


def test_checks_leave_other_accounts_balances_alone(client, household):
    """Money moved to other accounts is meant to be spent, so their balance
    drifting from the plan is not reported."""
    savings = account(household, "savings")
    add(client, household, "Alex", "income", "Salary", [30000] * 12)
    add(client, household, "Shared", "expense", "Rent", [20000] * 12)
    client.put(f"/api/accounts/{savings['id']}/balance", json={"year": 2027, "month": 1, "amount": 500000})
    client.put(f"/api/accounts/{savings['id']}/balance", json={"year": 2027, "month": 3, "amount": 5000})
    assert not any(f["type"] == "balance" for f in client.get("/api/checks").json()["findings"])


def test_one_budget_account_and_expense_accounts(client, household):
    main = account(household, "main")
    card = client.post("/api/accounts", json={"name": "Card", "kind": "expense"}).json()
    assert card["in_overview"] == 0  # expense accounts stay off the overview unless chosen
    # Making it the budget account turns the old one into an expense account.
    client.patch(f"/api/accounts/{card['id']}", json={"kind": "main"})
    kinds = {a["id"]: a["kind"] for a in client.get("/api/state").json()["accounts"]}
    assert kinds[card["id"]] == "main" and kinds[main["id"]] == "expense"
    assert list(kinds.values()).count("main") == 1
    # An expense account shows its real balances only, nothing worked out from the budget.
    client.put(f"/api/accounts/{main['id']}/balance", json={"year": 2027, "month": 2, "amount": 900})
    acc = next(a for a in client.get("/api/summary/2027").json()["accounts"] if a["id"] == main["id"])
    assert acc["balances"][1] == 900 and acc["balances"][2] is None


def test_migration_keeps_one_budget_account(tmp_path):
    import sqlite3
    conn = sqlite3.connect(tmp_path / "old.db")
    conn.row_factory = sqlite3.Row
    db.migrate(conn, db.MIGRATIONS[:5])  # as before expense accounts
    conn.executescript("""
        INSERT INTO accounts (id, name, kind, sort) VALUES (1, 'A', 'main', 0), (2, 'B', 'main', 1), (3, 'S', 'savings', 2);
        INSERT INTO account_balances (account_id, year, month, cents) VALUES (2, 2027, 1, 100);
    """)
    db.migrate(conn)
    rows = {r["id"]: (r["kind"], r["in_overview"]) for r in conn.execute("SELECT * FROM accounts")}
    assert rows == {1: ("main", 1), 2: ("expense", 0), 3: ("savings", 1)}
    assert conn.execute("SELECT COUNT(*) FROM account_balances").fetchone()[0] == 1  # nothing cascaded away


def test_measure_day():
    from datetime import date
    from app.banking import measure_day
    assert measure_day(2026, 10, "last") == date(2026, 10, 31)
    assert measure_day(2026, 10, "before_payday") == date(2026, 10, 29)  # last weekday Fri 30th
    assert measure_day(2026, 5, "before_payday") == date(2026, 5, 28)  # 31st is a Sunday: Friday 29th



def test_restore_converts_an_old_backup(client, household):
    """A backup from before expense accounts (and from when the app was called
    Jetbuddy) is brought up to date: one budget account, the rest expense accounts."""
    old = client.get("/api/export").json()
    old["format"], old["schema"] = "jetbuddy", 5
    for a in old["accounts"]:
        a.pop("in_overview")
    old["accounts"].append({**old["accounts"][0], "id": 99, "name": "Second", "sort": 9})
    assert client.post("/api/restore", json=old).status_code == 200
    kinds = {a["name"]: (a["kind"], a["in_overview"]) for a in client.get("/api/state").json()["accounts"]}
    first = old["accounts"][0]["name"]
    assert kinds[first] == ("main", True) and kinds["Second"] == ("expense", False)


def test_delete_a_budget_year(client, household):
    keep = add(client, household, "Alex", "income", "Salary", [31250] * 12)  # 2027
    client.post("/api/years/2028/copy", json={"from_year": 2027})
    only_2028 = client.post("/api/items", json={"member_id": member(household, "Alex")["id"], "kind": "expense",
                                                 "name": "New car", "year": 2028, "amounts": [500] * 12}).json()
    state = client.delete("/api/years/2028").json()
    assert 2028 not in state["years"] and 2027 in state["years"]
    assert client.get("/api/years/2028/items").json() == []
    assert [i["id"] for i in client.get("/api/years/2027/items").json()] == [keep["id"]]
    assert client.patch(f"/api/items/{only_2028['id']}?year=2028", json={"name": "x"}).status_code == 404
