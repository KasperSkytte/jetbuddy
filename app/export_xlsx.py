"""The budget as an Excel workbook, laid out for someone else to read as well
as yourself. Values, not formulas, so it reads
the same in every viewer (formulas show as empty in some previews)."""

import io
import json
import sqlite3
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import calc

TEXT = {
    "en": {
        "summary": "Summary", "budget": "Budget", "balances": "Account balances",
        "title": "Budget {year}", "made": "Made {date} with Sparegrisen. Amounts in {currency}.",
        "household": "Household", "key_figures": "Key figures",
        "per_month": "Per month (avg.)", "per_year": "Whole year",
        "income": "Income", "expenses": "Expenses", "total_income": "Total income", "total_expenses": "Total expenses",
        "of_fixed": "of which fixed expenses", "of_savings": "of which transferred to savings",
        "surplus": "Surplus (income − expenses)", "saved": "Saved in total (savings + surplus)", "savings_rate": "Savings rate",
        "after_fixed": "Disposable income (after fixed expenses)",
        "per_person": "Per person", "person": "Person", "income_month": "Income per month", "expenses_month": "Expenses per month",
        "income_year": "Income per year", "expenses_year": "Expenses per year",
        "item": "Item", "who": "Who", "fixed": "Fixed", "total": "Total", "avg": "Avg./month",
        "total_for": "Total {name}", "uncategorized": "Uncategorized",
        "accounts": "Account balances", "account": "Account", "latest_real": "Latest real balance", "at": "End of",
        "expected_dec": "Expected end of December",
        "month": "Month",
        "balances_note": "Black: real balance (from the bank or entered). Grey italics: expected from the budget.",
        "fixed_mark": "✓",
    },
    "da": {
        "summary": "Oversigt", "budget": "Budget", "balances": "Kontosaldi",
        "title": "Budget {year}", "made": "Udarbejdet {date} med Sparegrisen. Beløb i {currency}.",
        "household": "Husstand", "key_figures": "Nøgletal",
        "per_month": "Pr. måned (gns.)", "per_year": "Hele året",
        "income": "Indtægter", "expenses": "Udgifter", "total_income": "Indtægter i alt", "total_expenses": "Udgifter i alt",
        "of_fixed": "heraf faste udgifter", "of_savings": "heraf overført til opsparing",
        "surplus": "Overskud (indtægter − udgifter)", "saved": "Opsparing i alt (opsparing + overskud)", "savings_rate": "Opsparingsrate",
        "after_fixed": "Rådighedsbeløb",
        "per_person": "Pr. person", "person": "Person", "income_month": "Indtægter pr. md.", "expenses_month": "Udgifter pr. md.",
        "income_year": "Indtægter pr. år", "expenses_year": "Udgifter pr. år",
        "item": "Post", "who": "Hvem", "fixed": "Fast", "total": "I alt", "avg": "Gns./md.",
        "total_for": "I alt {name}", "uncategorized": "Uden kategori",
        "accounts": "Kontosaldi", "account": "Konto", "latest_real": "Seneste faktiske saldo", "at": "Ultimo",
        "expected_dec": "Forventet ultimo december",
        "month": "Måned",
        "balances_note": "Sort: faktisk saldo (fra banken eller indtastet). Grå kursiv: forventet ud fra budgettet.",
        "fixed_mark": "✓",
    },
}
MONTHS = {
    "en": ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
    "da": ["Jan", "Feb", "Mar", "Apr", "Maj", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dec"],
}
MONTHS_LONG = {
    "en": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"],
    "da": ["januar", "februar", "marts", "april", "maj", "juni", "juli", "august", "september", "oktober", "november", "december"],
}

MONEY = "#,##0;-#,##0"
MONEY_RED = "#,##0;[Red]-#,##0"
PCT = "0%"
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16)
SUBTLE = Font(color="666666", italic=True)
HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
GROUP_FILL = PatternFill("solid", fgColor="F3F3F1")
TOTAL_FILL = PatternFill("solid", fgColor="DDE6F3")
THIN = Side(style="thin", color="B0B0B0")
TOP = Border(top=THIN)
DOUBLE = Border(top=THIN, bottom=Side(style="double", color="808080"))


def _money(cents: int | None) -> float | None:
    return None if cents is None else cents / 100


def _setup_print(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False


def _row(ws, r, values, *, font=None, fill=None, fmt=None, border=None, start=1):
    for i, v in enumerate(values):
        c = ws.cell(row=r, column=start + i, value=v)
        if font:
            c.font = font
        if fill:
            c.fill = fill
        if border:
            c.border = border
        if fmt and isinstance(v, (int, float)) and not isinstance(v, bool):
            c.number_format = fmt


def build(conn: sqlite3.Connection, year: int, group: str = "auto") -> tuple[bytes, str]:
    settings = {r["key"]: json.loads(r["value"]) for r in conn.execute("SELECT key, value FROM settings")}
    lang = settings.get("language", "en")[:2]
    lang = lang if lang in TEXT else "en"
    T = TEXT[lang]
    months = MONTHS[lang]
    currency = settings.get("currency", "")
    household = settings.get("household", "")

    members = {r["id"]: dict(r) for r in conn.execute("SELECT * FROM members ORDER BY is_shared, sort, id")}
    categories = {r["id"]: dict(r) for r in conn.execute("SELECT * FROM categories ORDER BY sort, id")}
    accounts = {r["id"]: dict(r) for r in conn.execute("SELECT * FROM accounts WHERE archived = 0 ORDER BY sort, id")}
    ids = calc.visible_item_ids(conn, year)
    items = (
        [dict(r) for r in conn.execute(f"SELECT * FROM items WHERE id IN ({','.join('?' * len(ids))}) ORDER BY sort, id", ids)]
        if ids else []
    )
    amounts = calc.year_amounts(conn, year)
    summary = calc.summary(conn, year)

    for item in items:
        item["values"] = amounts.get(item["id"], [None] * 12)
        item["total"] = sum(v or 0 for v in item["values"])

    if group == "auto":
        # By category when the expenses are mostly categorised, else by person.
        exp = [i for i in items if i["kind"] == "expense"]
        categorised = sum(i["total"] for i in exp if i["category_id"])
        group = "category" if exp and categorised >= 0.5 * sum(i["total"] for i in exp) else "member"

    income = summary["income"]
    expense = summary["expense"]
    months_with_data = sum(1 for i in range(12) if income[i] or expense[i]) or 12
    per_month = lambda cents: _money(round(cents / months_with_data))

    wb = Workbook()

    # ---------------------------------------------------------------- summary
    ws = wb.active
    ws.title = T["summary"]
    ws["A1"] = T["title"].format(year=year)
    ws["A1"].font = TITLE
    ws["A2"] = household
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = T["made"].format(date=_long_date(date.today(), lang), currency=currency)
    ws["A3"].font = SUBTLE
    people = [m["name"] for m in members.values() if not m["is_shared"] and not m["archived"]]
    ws["A5"], ws["B5"] = T["household"], ", ".join(people)
    ws["A5"].font = BOLD

    r = 7
    _row(ws, r, [T["key_figures"], T["per_month"], T["per_year"]], font=BOLD, fill=HEADER_FILL)
    total_in, total_out = sum(income), sum(expense)
    fixed, savings = sum(summary["fixed"]), sum(summary["savings"])
    net = total_in - total_out
    # Expenses are negative, as everywhere in the app.
    lines = [(T["total_income"], total_in, False), (T["total_expenses"], -total_out, False)]
    if fixed:
        lines.append(("   " + T["of_fixed"], -fixed, False))
    if savings:
        lines.append(("   " + T["of_savings"], -savings, False))
    lines.append((T["surplus"], net, True))
    if fixed:
        lines.append((T["after_fixed"], total_in - fixed, False))
    if savings:
        lines.append((T["saved"], savings + net, True))
    for label, cents, strong in lines:
        r += 1
        _row(ws, r, [label, per_month(cents), _money(cents)], fmt=MONEY_RED if strong else MONEY, font=BOLD if strong else None)
    if total_in:
        r += 1
        _row(ws, r, [T["savings_rate"], None, (savings + net) / total_in], fmt=PCT)

    r += 2
    _row(ws, r, [T["per_person"], T["income_month"], T["expenses_month"], T["income_year"], T["expenses_year"]], font=BOLD, fill=HEADER_FILL)
    by_member = {m["id"]: m for m in summary["members"]}
    for m in members.values():
        data = by_member.get(m["id"])
        if not data:
            continue
        mi, me = sum(data["income"]), sum(data["expense"])
        if not (mi or me):
            continue
        r += 1
        _row(ws, r, [m["name"], per_month(mi), per_month(-me), _money(mi), _money(-me)], fmt=MONEY)

    # Account balances: the latest real one, and where the budget says the year ends.
    real_accounts = [a for a in summary["accounts"] if a["id"] in accounts and any(a["manual"])]
    if real_accounts:
        r += 2
        _row(ws, r, [T["accounts"], T["latest_real"], "", T["expected_dec"]], font=BOLD, fill=HEADER_FILL)
        for a in real_accounts:
            last = max(i for i in range(12) if a["manual"][i])
            r += 1
            _row(ws, r, [accounts[a["id"]]["name"], _money(a["balances"][last]), f"{T['at']} {MONTHS_LONG[lang][last]}", _money(a["balances"][11])], fmt=MONEY)

    ws.column_dimensions["A"].width = 42
    for col in "BCDE":
        ws.column_dimensions[col].width = 20
    _setup_print(ws, landscape=False)

    # ---------------------------------------------------------------- budget
    ws = wb.create_sheet(T["budget"])
    ws["A1"] = f"{T['title'].format(year=year)} · {household}"
    ws["A1"].font = TITLE
    header = [T["item"], T["who"], T["fixed"], *months, T["total"], T["avg"]]
    _row(ws, 3, header, font=BOLD, fill=HEADER_FILL)
    first_money_col = 4
    r = 3

    # Expenses are written as negative numbers, as the app shows them.
    def sign(kind):
        return -1 if kind == "expense" else 1

    def write_item(item):
        nonlocal r
        r += 1
        k = sign(item["kind"])
        values = [_money(None if v is None else k * v) for v in item["values"]]
        _row(ws, r, [item["name"], members.get(item["member_id"], {}).get("name", ""), T["fixed_mark"] if item["fixed"] else "",
                     *values, _money(k * item["total"]), _money(round(k * item["total"] / 12))], fmt=MONEY)

    def write_total(label, group_items, style="sub"):
        nonlocal r
        r += 1
        month_totals = [sum(sign(i["kind"]) * (i["values"][m] or 0) for i in group_items) for m in range(12)]
        total = sum(month_totals)
        _row(ws, r, [label, "", "", *[_money(v) for v in month_totals], _money(total), _money(round(total / 12))],
             fmt=MONEY, font=BOLD, fill=TOTAL_FILL if style == "total" else None, border=DOUBLE if style == "total" else TOP)
        return month_totals

    def write_group(name, group_items):
        nonlocal r
        r += 1
        _row(ws, r, [name] + [""] * (len(header) - 1), font=BOLD, fill=GROUP_FILL)
        for item in group_items:
            write_item(item)
        if len(group_items) > 1:
            write_total(T["total_for"].format(name=name), group_items)

    def groups_of(kind, by):
        sel = [i for i in items if i["kind"] == kind]
        out = defaultdict(list)
        for i in sel:
            out[i["member_id"] if by == "member" else i["category_id"]].append(i)
        if by == "member":
            order = [mid for mid in members if mid in out]
            return [(members[mid]["name"], out[mid]) for mid in order], sel
        order = [cid for cid in categories if cid in out] + ([None] if None in out else [])
        return [(categories[cid]["name"] if cid else T["uncategorized"], out[cid]) for cid in order], sel

    r += 1
    ws.cell(row=r, column=1, value=T["income"].upper()).font = Font(bold=True, color="2A78D6")
    inc_groups, inc_items = groups_of("income", "member")
    for name, g in inc_groups:
        write_group(name, g)
    inc_totals = write_total(T["total_income"], inc_items, "total")

    r += 2
    ws.cell(row=r, column=1, value=T["expenses"].upper()).font = Font(bold=True, color="C2410C")
    exp_groups, exp_items = groups_of("expense", group)
    for name, g in exp_groups:
        write_group(name, g)
    exp_totals = write_total(T["total_expenses"], exp_items, "total")

    r += 2
    surplus = [inc_totals[m] + exp_totals[m] for m in range(12)]  # the expenses are negative
    _row(ws, r, [T["surplus"], "", "", *[_money(v) for v in surplus], _money(sum(surplus)), _money(round(sum(surplus) / 12))],
         fmt=MONEY_RED, font=BOLD, fill=TOTAL_FILL, border=DOUBLE)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 6
    for col in range(first_money_col, first_money_col + 14):
        ws.column_dimensions[get_column_letter(col)].width = 10.5
    ws.freeze_panes = "D4"
    for row in ws.iter_rows(min_row=3, max_row=3):
        for c in row[2:]:
            c.alignment = Alignment(horizontal="right")
    _setup_print(ws)

    # ---------------------------------------------------------------- account balances
    acc_rows = [a for a in summary["accounts"] if a["id"] in accounts and any(v is not None for v in a["balances"])]
    if acc_rows:
        ws = wb.create_sheet(T["balances"])
        ws["A1"] = f"{T['balances']} {year} · {household}"
        ws["A1"].font = TITLE
        ws["A2"] = T["balances_note"]
        ws["A2"].font = SUBTLE
        _row(ws, 4, [T["month"], *[accounts[a["id"]]["name"] for a in acc_rows]], font=BOLD, fill=HEADER_FILL)
        for m in range(12):
            ws.cell(row=5 + m, column=1, value=MONTHS_LONG[lang][m].capitalize())
            for j, a in enumerate(acc_rows):
                v = a["balances"][m]
                c = ws.cell(row=5 + m, column=2 + j, value=_money(v))
                c.number_format = MONEY_RED
                if not a["manual"][m]:
                    c.font = Font(color="808080", italic=True)
        ws.column_dimensions["A"].width = 16
        for j in range(len(acc_rows)):
            ws.column_dimensions[get_column_letter(2 + j)].width = 18
        _setup_print(ws, landscape=False)

    buf = io.BytesIO()
    wb.save(buf)
    safe = "".join(ch for ch in household if ch.isalnum() or ch in " -_").strip().replace(" ", "-") or "budget"
    return buf.getvalue(), f"{T['budget']}-{year}-{safe}.xlsx"


def _long_date(d: date, lang: str) -> str:
    if lang == "da":
        return f"{d.day}. {MONTHS_LONG['da'][d.month - 1]} {d.year}"
    return f"{MONTHS_LONG['en'][d.month - 1]} {d.day}, {d.year}"
